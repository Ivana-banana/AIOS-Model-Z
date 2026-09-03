"""Валидационный кейс приходит с чужими условиями — проверка подстановкой.

ТЗ п. 4.3 и 5.1: в рейтинг идёт ЧДД на кейсе организаторов, а не наше число на
Модели Z. Меняются при этом **условия задачи, а не объект**: нормативы, пределы
объекта, окно управления, состав фонда. Значит, ценность проверки не в
рассуждении «у нас всё в конфиге», а в подстановке: меняем величину, запускаем,
смотрим вывод.

Каждый тест здесь закрывает один дефект, найденный такой подстановкой 23.08
(задача Б28). Разбор целиком — `docs/validation-case.md`.
"""

from __future__ import annotations

import ast
import csv
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from conftest import make_row
from openpyxl import load_workbook

from aios import economics, loop, summary
from aios.margin import Norms, screen

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def alien_norms(tmp_path: Path) -> Path:
    """Нормативы организаторов «как на валидации»: цена нефти 42 000 вместо 28 000."""
    path = tmp_path / "Нормативы_организаторов.xlsx"
    shutil.copy(economics.NORMS_FILE, path)
    book = load_workbook(path)
    for row in book["Нормативы"].iter_rows(min_row=2):
        if row[0].value == "oilPriceRubT":
            row[2].value = 42000.0
    book.save(path)
    return path


@pytest.fixture
def alien_pumps(tmp_path: Path) -> Path:
    """Нормативы, где старший ЭЦН закрывает 250–400, а не 320–500."""
    path = tmp_path / "Нормативы_с_другим_ЭЦН.xlsx"
    shutil.copy(economics.NORMS_FILE, path)
    book = load_workbook(path)
    sheet = book["ЭЦН"]
    for row in sheet.iter_rows(min_row=2):
        if row[0].value in (400.0, 400):
            sheet.delete_rows(row[0].row)
            break
    book.save(path)
    return path


def rows_for(well: str, watercut: float, months: int = 12) -> list[dict[str, object]]:
    fluids = summary.Fluids()
    rows = []
    for index in range(months):
        liquid_m3 = 30.0 * 30
        oil_m3 = liquid_m3 * (1 - watercut)
        oil_t = oil_m3 * fluids.oil_density / 1000
        liquid_t = oil_t + (liquid_m3 - oil_m3) * fluids.water_density / 1000
        rows.append(make_row(date(2014, index + 1, 1).isoformat(), well, wlpr=30.0,
                             womr=oil_t / 30, liquid=liquid_t, oil=oil_t))
    return rows


def test_alien_norms_change_the_declared_number(alien_norms: Path):
    """Другой файл нормативов обязан менять ЧДД, а не только маржу скрининга.

    Замер 23.08 до правки: на странице прогона маржа скважины шла по файлу из
    `[economics].norms_file` (53,4 → 148,1 млн при цене 42 000), а ЧДД считался
    по вендоренному файлу и оставался 11,8 млн. Одна страница, две разные цены
    нефти, и увидеть это можно было только в исходниках.
    """
    rows = rows_for("A", 0.5)
    ours = economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                                 start_date="2014-01-01")
    theirs = economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                                   start_date="2014-01-01", norms_file=alien_norms)

    assert theirs.chdd_m > ours.chdd_m * 1.5
    # скрининг и ЧДД слышат один и тот же файл
    assert screen(rows, Norms.from_file(alien_norms))[0].total_margin_m > \
        screen(rows, Norms.from_file())[0].total_margin_m


def test_the_cycle_confirms_by_the_norms_it_was_given(alien_norms: Path):
    """`loop.confirm` — последний шаг перед заявкой числа, и нормативы у него свои."""
    rows = rows_for("A", 0.5)
    result = loop.CycleResult(
        base_chdd_m=0.0, chdd_m=0.0, actions=(), rate_choices=(), allocations=(),
        steps=(), warnings=(), problems=(),
    )

    ours = loop.confirm(result, rows, date(2014, 1, 1), base_rows=rows)
    theirs = loop.confirm(result, rows, date(2014, 1, 1), base_rows=rows,
                          norms_file=alien_norms)

    assert theirs.chdd_m != ours.chdd_m
    assert theirs.base_chdd_m != ours.base_chdd_m


def test_the_web_page_shows_the_number_of_the_configured_norms(tmp_path: Path,
                                                               alien_norms: Path):
    """UI — тоже путь, по которому число попадает человеку в глаза."""
    from fastapi.testclient import TestClient

    from aios.web.app import create_app

    config = (ROOT / "tests" / "test_web.py").read_text(encoding="utf-8")
    config = config.split('CONFIG = """')[1].split('"""')[0]
    project = tmp_path / "project"
    run_dir = project / "runs" / "base"
    run_dir.mkdir(parents=True)
    (project / "config.toml").write_text(
        config.format(norms=alien_norms.resolve().as_posix()), encoding="utf-8"
    )
    (run_dir / "wells_schedule.inc").write_text("DATES\n 01 JAN 2014 /\n/\n", encoding="utf-8")
    rows = rows_for("A", 0.5)
    with (run_dir / "rows.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=economics.REQUIRED_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    page = TestClient(create_app(project)).get("/run/base").text

    expected = economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                                     start_date="2014-01-01",
                                     norms_file=alien_norms).chdd_m
    assert f"{expected:.1f}".replace(".", ",") in page or f"{expected:.1f}" in page


def _our_sources() -> list[Path]:
    """Наши модули и скрипты — списком из git, а не маской по каталогу.

    На Windows каталог venv `Scripts/` — для файловой системы тот же путь, что
    наш `scripts/`, поэтому маска затянула бы чужие файлы.
    """
    listed = subprocess.run(
        ["git", "ls-files", "aios", "scripts"],
        cwd=ROOT, capture_output=True, text=True, check=True,
    ).stdout.split()
    return [ROOT / name for name in listed if name.endswith(".py")]


def test_every_npv_call_names_its_norms():
    """Ни один вызов расчётчика в нашем коде не берёт нормативы молча.

    Умолчание `compute_npv` — файл действующей поставки, и это верно для тестов
    и для Модели Z. Но в рабочем пути умолчание означает «кейс организаторов
    посчитан по нашим нормативам», а увидеть это нельзя ниоткуда: ошибки не
    будет, будет другое число. Проверка статическая — она не требует ни
    симулятора, ни кейса.
    """
    offenders: list[str] = []
    for path in _our_sources():
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            name = getattr(node.func, "attr", None) or getattr(node.func, "id", None)
            if name not in {"compute_npv", "compute_npv_from_file"}:
                continue
            passed = {keyword.arg for keyword in node.keywords}
            if not passed & {"norms_file", "assumptions"}:
                offenders.append(f"{path.relative_to(ROOT).as_posix()}:{node.lineno}")

    assert offenders == [], (
        "вызовы ЧДД без явного источника нормативов:\n  " + "\n  ".join(offenders)
    )


def test_the_liquid_limit_never_exceeds_the_vendor_ceiling(tmp_path: Path):
    """Таблица ЭЦН кейса не может разрешить больше, чем принимает расчётчик.

    У эталона потолок зашит литералом: `if row["WLPR"] > 500` в
    `compute_calculation`. Пока предел брался как «правая граница старшего
    насоса», таблица с насосом на 600 заставила бы наш валидатор пропустить 550,
    а расчётчик организаторов отверг бы этот же файл с `ValueError` — на том
    единственном прогоне, который оценивают (ревизия 03.09, 2.3 #4).
    """
    import shutil

    from openpyxl import load_workbook

    from aios import schedule as schedule_module

    path = tmp_path / "Нормативы_с_большим_ЭЦН.xlsx"
    shutil.copy(economics.NORMS_FILE, path)
    book = load_workbook(path)
    sheet = book["ЭЦН"]
    for row in sheet.iter_rows(min_row=2):
        if row[2].value in (500.0, 500):
            row[2].value = 600.0
            break
    book.save(path)

    _assumptions, pumps = economics.load_norms(path)
    assert max(float(pump["max"]) for pump in pumps) == 600.0
    assert schedule_module.liquid_limit(pumps) == 500.0

    rows = [make_row(f"2014-{month:02d}-01", "A", wlpr=550.0, womr=90.0,
                     liquid=16500.0, oil=2700.0) for month in (1, 2, 3)]
    with pytest.raises(ValueError, match="500"):
        economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                              start_date="2014-01-01", norms_file=path)


def test_the_injection_pressure_limit_comes_from_the_case(tmp_path: Path):
    """Предел давления нагнетания — из `[limits]`, а не литерал 300 от Модели Z.

    Кейс со своим пределом ловил только планировщик; валидатор расписания и
    предполётная проверка держали литерал, то есть кейс с потолком 250 атм
    прошёл бы обе (ревизия 03.09, 2.3 #5).
    """
    from aios import schedule as schedule_module
    from aios.deck import InjectorControl, ScheduleStep

    steps = [ScheduleStep(date(2014, 1, 1), (),
                          (InjectorControl("I1", "WATER", "OPEN", "RATE", 200.0, 280.0),))]

    assert schedule_module.validate(steps) == []
    problems = schedule_module.validate(steps, max_injection_bhp=250.0)
    assert len(problems) == 1 and "280" in str(problems[0]) and "250" in str(problems[0])


def test_an_unreadable_start_date_is_refused_not_replaced():
    """Нераспознанную дату эталон молча меняет на первую дату данных.

    `requested_start = normalize_date(start_date) or min_date` — то есть `"2010"`
    не ошибка, а тихий возврат к умолчанию: базовый год уезжает, ЧДД меняется,
    предупреждений нет (ревизия 03.09, 2.3 #3).
    """
    rows = [make_row(f"2014-{month:02d}-01", "A", wlpr=30.0, womr=15.0,
                     liquid=900.0, oil=450.0) for month in (1, 2, 3)]

    with pytest.raises(ValueError, match="не распознано"):
        economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                              start_date="2010", norms_file=economics.NORMS_FILE)
    with pytest.raises(ValueError, match="не распознано"):
        economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                              start_date="01 JAN 2010", norms_file=economics.NORMS_FILE)

    # осознанное умолчание эталона по-прежнему законно и видно в коде вызова
    assert economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                                 start_date=None,
                                 norms_file=economics.NORMS_FILE).chdd_m != 0.0


def test_the_liquid_limit_travels_with_the_pump_table(alien_pumps: Path):
    """Предел дебита жидкости — из таблицы ЭЦН кейса, а не литерал 500 от Модели Z.

    Методика запрещает дебит выше верхней границы старшего насоса. Число 500
    принадлежит `Нормативы_ЧДД.xlsx` действующей поставки; на валидационном
    кейсе (ТЗ п. 4.3) таблица приходит от организаторов. Литерал в коде означал
    бы, что мы либо генерируем режим, который их расчётчик отвергнет, либо
    запрещаем законный.
    """
    from aios import schedule as schedule_module
    from aios.deck import ProducerControl, ScheduleStep

    _assumptions, pumps = economics.load_norms(alien_pumps)
    limit = schedule_module.liquid_limit(pumps)
    assert limit == 400.0
    assert schedule_module.liquid_limit(economics.load_norms()[1]) ==         schedule_module.MAX_LIQUID_RATE_M3_PER_DAY

    steps = [ScheduleStep(date(2014, 1, 1),
                          (ProducerControl("A", "OPEN", "LRAT", 450.0, 50.0),), ())]
    assert schedule_module.validate(steps) == []
    problems = schedule_module.validate(steps, max_liquid_rate=limit)
    assert len(problems) == 1 and "450" in str(problems[0])


def test_a_non_monthly_export_is_refused_by_name():
    """Строка = скважино-месяц: квартальный шаг молча удешевляет фонд втрое.

    Замер 23.08: тот же годовой объём на квартальном шаге даёт OPEX фонда 0,333
    вместо 1,0 млн руб. в год и ЧДД на 1,7 % выше — ни ошибки, ни предупреждения
    при этом не возникало.
    """
    monthly = [make_row(f"2014-{month:02d}-01", "A", wlpr=30.0, womr=5.0,
                        liquid=900.0, oil=450.0) for month in range(1, 13)]
    quarterly = [row for row in monthly if int(str(row["DATA"])[5:7]) in (1, 4, 7, 10)]

    assert summary.sanity_problems(monthly) == []
    problems = summary.sanity_problems(quarterly)
    assert len(problems) == 1
    assert "не месячный" in problems[0]

    # и деньги действительно разъезжаются — ради этого проверка и стоит
    money = [economics.compute_npv(rows, headers=economics.REQUIRED_COLUMNS,
                                   start_date="2014-01-01",
                                   norms_file=economics.NORMS_FILE).annual[0]["fundOpexM"]
             for rows in (monthly, quarterly)]
    assert money[0] == pytest.approx(1.0)
    assert money[1] == pytest.approx(1 / 3)


def test_planner_reports_a_fractional_limit_as_it_is():
    """Предел объекта печатается без округления: «предела 0» при пределе 0,5 — ложь."""
    from aios import planner
    from aios.deck import ProducerControl, ScheduleStep

    steps = [ScheduleStep(date(2014, 1, 1),
                          (ProducerControl("A", "OPEN", "LRAT", 30.0, 50.0),), ())]
    forecast = [make_row("2014-01-01", "A", wlpr=30.0, womr=5.0, liquid=900.0, oil=450.0)]
    limits = planner.Limits(max_field_liquid_m3_per_day=0.5, enforce_water_balance=False)

    _fitted, decisions = planner.enforce(steps, forecast, Norms.from_file(), limits)

    assert decisions and "предела 0,5" in decisions[0].reason.replace(".", ",")


# --- Б29: чужие режимы управления -------------------------------------------
#
# `optimise.scale` умножает поле расхода независимо от режима. Под `LRAT` и
# `RATE` это уставка, и множитель осмыслен; под `BHP`, `THP`, `GRUP`, `ORAT`,
# `WRAT`, `RESV` то же поле — верхний предел. На Модели Z все 13 382 записи
# добывающих под `LRAT` и все 8 212 записей нагнетательных под `RATE` (замер
# 23.08), поэтому путь не исполнялся ни разу.

WELLS_HEAD = """WELSPECS
 'P1' 'G' 1 1 1* 'OIL' /
 'I1' 'G' 2 2 1* 'WATER' /
/

DATES
 01 JAN 2007 /
/

"""


def _deck_project(tmp_path: Path, schedule_text: str) -> Path:
    """Крошечный дек с заданной SCHEDULE-секцией и конфиг под него."""
    from test_dataset import DECK_DATA, TINY_CONFIG

    deck = tmp_path / "deck"
    deck.mkdir()
    (deck / "FIELD.DATA").write_text(DECK_DATA, encoding="utf-8")
    for name in ("p.inc", "r.inc"):
        (deck / name).write_text("-- пусто\n", encoding="utf-8")
    (deck / "h.inc").write_text(schedule_text, encoding="utf-8")
    config = tmp_path / "config.toml"
    config.write_text(
        TINY_CONFIG.format(norms=Path(economics.NORMS_FILE).resolve().as_posix()),
        encoding="utf-8",
    )
    return config


@pytest.mark.parametrize("mode, record", [
    ("ORAT", " 'P1' 'OPEN' 'ORAT' 20.0 1* 1* 1* 1* 50 1* 1* /"),
    ("WRAT", " 'P1' 'OPEN' 'WRAT' 1* 25.0 1* 1* 1* 50 1* 1* /"),
    ("RESV", " 'P1' 'OPEN' 'RESV' 1* 1* 1* 1* 40.0 50 1* 1* /"),
])
def test_a_producer_without_a_rate_field_is_not_shut_by_scaling(mode: str, record: str):
    """Работающая скважина без колонки LRAT не должна закрываться множителем.

    До 23.08 `_scaled` считал `(None or 0) × 1.5 = 0` и превращал запись в
    `SHUT`: на чужом деке сценарий «форсаж ×1.5» **останавливал** скважины под
    `ORAT`, `WRAT` и `RESV`, а печатная нота при этом говорила «поле дебита это
    предел, а не уставка» — ровно наоборот произошедшему.
    """
    from aios.deck import parse_schedule_text
    from aios.optimise import scale

    steps = parse_schedule_text(WELLS_HEAD + "WCONPROD\n" + record + "\n/\n").steps
    scaled = scale(steps, 1.5)

    control = scaled[0].producers[0]
    assert control.status == "OPEN", f"{mode}: скважина закрыта множителем"
    assert control.mode == mode
    assert control.liquid_rate is None, "поля дебита не было — не должно и появиться"


def test_an_injector_without_a_rate_field_is_not_shut_by_scaling():
    """То же для нагнетательной под `RESV`: приёмистость стоит в своей колонке."""
    from aios.deck import parse_schedule_text
    from aios.optimise import scale_injection

    text = WELLS_HEAD + "WCONINJE\n 'I1' 'WATER' 'OPEN' 'RESV' 1* 220.0 300 1* /\n/\n"
    scaled = scale_injection(parse_schedule_text(text).steps, 1.6)

    control = scaled[0].injectors[0]
    assert control.status == "OPEN"
    assert control.rate is None


def test_the_scaling_report_classifies_every_mode():
    """Каждый режим попадает ровно в одну корзину, и доля считается по всем.

    Раньше нота смотрела только на добывающих: нагнетательная под `RESV` молча
    останавливалась, а строка сообщала «множитель применён буквально ко всем
    записям, все под LRAT».
    """
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep
    from aios.optimise import scale, scale_injection, scaling_report

    before = [ScheduleStep(date(2007, 1, 1), (
        ProducerControl("уставка", "OPEN", "LRAT", 100.0, 50.0),
        ProducerControl("предел", "OPEN", "LRAT", 400.0, 50.0),
        ProducerControl("давление", "OPEN", "BHP", 100.0, 90.0),
        ProducerControl("устье", "OPEN", "THP", 100.0, 90.0),
        ProducerControl("группа", "OPEN", "GRUP", 100.0, 50.0),
        ProducerControl("нефть", "OPEN", "ORAT", None, 50.0),
        ProducerControl("стоит", "SHUT", "LRAT", None, 50.0),
    ), (
        InjectorControl("расход", "WATER", "OPEN", "RATE", 100.0, 300.0),
        InjectorControl("давление", "WATER", "OPEN", "BHP", 100.0, 300.0),
        InjectorControl("объём", "WATER", "OPEN", "RESV", None, 300.0),
    ))]
    after = scale_injection(scale(before, 1.5), 1.6)
    report = scaling_report(before, after, production=1.5, injection=1.6)

    assert report.records == 9, "остановленная запись не масштабируется и не считается"
    assert report.literal == 2, "буквально применён только к LRAT и RATE с уставкой"
    assert report.capped == 1
    assert report.off_mode == {"BHP": 2, "GRUP": 1, "THP": 1}
    assert report.without_rate == {"ORAT": 1, "RESV": 1}
    assert report.share == pytest.approx(2 / 9)


def test_the_preflight_refuses_a_label_that_does_not_describe_the_file(tmp_path: Path):
    """Доля буквально применённых записей — величина для машины, а не для лога."""
    import json
    import sys

    sys.path.insert(0, str(ROOT / "scripts"))
    from check_package import _scaling_problem

    schedule = tmp_path / "wells_schedule.inc"
    schedule.write_text("-- пусто\n", encoding="utf-8")

    missing, detail = _scaling_problem(schedule)
    assert missing is None and "не измерена" in detail

    (tmp_path / "scaling.json").write_text(
        json.dumps({"records": 10, "literal": 4, "share": 0.4}), encoding="utf-8"
    )
    problem, _detail = _scaling_problem(schedule)
    assert problem is not None and "40.0%" in problem.replace(",", ".")

    (tmp_path / "scaling.json").write_text(
        json.dumps({"records": 10, "literal": 10, "share": 1.0}), encoding="utf-8"
    )
    problem, detail = _scaling_problem(schedule)
    assert problem is None and "всем 10 записям" in detail


def test_a_deck_with_wconhist_is_refused_by_name(tmp_path: Path, capsys):
    """Дек с историческими режимами мы не считаем — и говорим это строкой.

    `WCONHIST` занимает место нашего `WCONPROD`: разобранные шаги остаются без
    режимов, а в собранный файл попали бы оба ключевых слова сразу. Отказ
    законен, трассировка на середине сборки — нет.
    """
    import sys

    sys.path.insert(0, str(ROOT / "scripts"))
    import build_scaled_case

    text = (WELLS_HEAD + "WCONHIST\n 'P1' 'OPEN' 'LRAT' 20.0 12.0 1* 1* 1* 1* /\n/\n"
            + "WCONINJH\n 'I1' 'WATER' 'OPEN' 200.0 1* 300 /\n/\n")
    config = _deck_project(tmp_path, text)

    code = build_scaled_case.main([str(tmp_path / "run"), "--production", "1.5",
                                   "--config", str(config)])

    assert code == 1
    printed = capsys.readouterr()
    assert "кейс не собран" in printed.err
    assert "WCONHIST" in printed.err


def test_the_matched_rule_never_asks_more_than_the_uniform_multiplier():
    """Правило «просить по достигнутому» обязано только снижать просьбу.

    На этом стоит весь довод защиты: уставка ниже пропускной способности даёт
    один и тот же расход в любом симуляторе, поэтому правило может оставить
    деньги на столе, но не завысить заявку. Свойство проверяется на доле выше
    единицы — скважина, взявшая больше запрошенного, не должна получить
    множитель выше общего.
    """
    import sys

    sys.path.insert(0, str(ROOT / "scripts"))
    from build_matched_case import multipliers

    wells = ["ровно", "половина", "сверх", "без замера"]
    ratios = {"ровно": 1.0, "половина": 0.5, "сверх": 1.4}

    chosen = multipliers(ratios, wells, injection=1.6, slack=1.0)

    assert max(chosen.values()) <= 1.6 + 1e-12
    assert chosen["половина"] == pytest.approx(0.8)
    assert chosen["сверх"] == pytest.approx(1.6), "доля выше единицы не поднимает просьбу"
    assert chosen["без замера"] == pytest.approx(1.6), "про скважину без замера ничего не известно"


def test_a_zero_delivered_ratio_is_named_and_not_a_silent_shutdown(tmp_path: Path):
    """Доля дошедшего, равная нулю, превращает масштабирование в остановку.

    Замер 23.08 (Б30): множитель 0 уходит в `optimise.scale_injection`, тот
    закрывает скважину, а отчёт масштабирования считает такую запись буквально
    применённым множителем — 7 из 7, 100 %. То есть ни одна проверка не увидела
    бы, что правило «просить меньше» кого-то остановило. Остановка стоит 1,0 млн
    руб. по методике и меняет энергию пласта: это отдельное решение, а не
    следствие арифметики.
    """
    import sys
    from datetime import date as date_type

    sys.path.insert(0, str(ROOT / "scripts"))
    from build_matched_case import multipliers

    from aios.deck import InjectorControl, ScheduleStep
    from aios.optimise import scale_injection, scaling_report

    chosen = multipliers({"пусто": 0.0, "живая": 1.0}, ["пусто", "живая"],
                         injection=1.6, slack=1.0)
    assert chosen["пусто"] == 0.0

    steps = [ScheduleStep(date_type(2007, 1, 1), (), (
        InjectorControl("пусто", "WATER", "OPEN", "RATE", 200.0, 300.0),
        InjectorControl("живая", "WATER", "OPEN", "RATE", 200.0, 300.0),
    ))]
    scaled = steps
    for well, multiplier in chosen.items():
        scaled = scale_injection(scaled, multiplier, {well})

    shut = [control for control in scaled[0].injectors if control.status == "SHUT"]
    assert [control.well for control in shut] == ["пусто"]

    # и вот почему на это нужна отдельная строка предупреждения: шлюз Р-92 такую
    # запись не отличает от честно применённого множителя
    report = scaling_report(steps, scaled, production={}, injection=chosen)
    assert report.share == 1.0, "отчёт масштабирования остановку не видит — это его предел"


def test_the_matched_case_names_the_missing_export_instead_of_a_traceback(tmp_path: Path,
                                                                          capsys):
    """Команда 4 инструкции зависит от вывода команды 3 — и обязана это сказать.

    Проверка 23.08 (Б30, раздел 3 `docs/submission.md`): без выгрузки первого
    расчёта команда уходила в трассировку pandas на двенадцать кадров. Это
    единственное, чего ей может не хватать на машине проверяющего.
    """
    import sys

    sys.path.insert(0, str(ROOT / "scripts"))
    import build_matched_case

    config = _deck_project(tmp_path, WELLS_HEAD
                           + "WCONPROD\n 'P1' 'OPEN' 'LRAT' 1* 1* 1* 30.0 1* 50 1* 1* /\n/\n"
                           + "WCONINJE\n 'I1' 'WATER' 'OPEN' 'RATE' 200.0 1* 300 1* /\n/\n")

    code = build_matched_case.main([str(tmp_path / "run"), "--rows",
                                    str(tmp_path / "нет-такой.csv"),
                                    "--production", "1.5", "--injection", "1.6",
                                    "--slack", "1.0", "--config", str(config)])

    assert code == 1
    printed = capsys.readouterr().err
    assert "нет выгрузки" in printed
    assert "export_rows.py" in printed, "отказ обязан называть, чем выгрузка получается"

# --- Б40: упреждающие ручки под ограничения кейса ---------------------------
#
# Организаторы сказали, что ограничения придут «по давлениям, по фонду, по
# закачке» (`docs/method/qa-organisers.md`). Каждая ручка проверяется тем же
# методом, что и вся Б28: подставили число — увидели, что срезано ровно на нём.
# И отдельно проверяется, что без ключа поведение не меняется: иначе новая ручка
# сдвинула бы заявляемое число.


def _plan(producers=((30.0, 50.0),), injectors=((200.0, 300.0),), months=3):
    """Расписание из нескольких шагов: «(расход, давление)» на скважину."""
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep

    steps = []
    for index in range(months):
        steps.append(ScheduleStep(
            date(2014, index + 1, 1),
            tuple(ProducerControl(f"P{n}", "OPEN", "LRAT", rate, bhp)
                  for n, (rate, bhp) in enumerate(producers, start=1)),
            tuple(InjectorControl(f"I{n}", "WATER", "OPEN", "RATE", rate, bhp)
                  for n, (rate, bhp) in enumerate(injectors, start=1)),
        ))
    return steps


def _forecast(wells=("P1",), months=3, liquid=900.0, oil=450.0):
    return [make_row(f"2014-{month:02d}-01", well, wlpr=30.0, womr=oil / 30,
                     liquid=liquid, oil=oil)
            for month in range(1, months + 1) for well in wells]


def _limits(**kwargs):
    from aios.planner import Limits

    base = {"enforce_water_balance": False}
    base.update(kwargs)
    return Limits(**base)


def test_pressure_limits_come_from_the_config_and_say_so():
    """Кейс «по давлениям» назовёт свои числа, а дек подменять нельзя.

    В деке Модели Z стоят 50 атм снизу у добывающих и 300 сверху у
    нагнетательных, и 84,4 % месяцев добычи идут ровно на нижнем. Предел из
    конфигурации перекрывает дек, и решение прямо говорит, что число взято
    оттуда: иначе на защите не отличить «так в деке» от «так мы задали».
    """
    from aios import planner

    steps = _plan(producers=((30.0, 50.0),), injectors=((200.0, 300.0),))

    fitted, decisions = planner.enforce(
        steps, _forecast(), Norms.from_file(),
        _limits(min_producer_bhp_atm=80.0, max_injector_bhp_atm=250.0),
    )

    assert fitted[0].producers[0].bhp == 80.0
    assert fitted[0].injectors[0].bhp == 250.0
    assert any("предел из конфига" in d.resource and "конфигурацией" in d.reason
               for d in decisions)

    # без ключей — ни одного решения и ни одного изменения
    untouched, none = planner.enforce(steps, _forecast(), Norms.from_file(), _limits())
    assert none == [] and untouched[0].producers[0].bhp == 50.0


def test_a_per_well_cap_cuts_the_well_the_field_limit_would_not():
    """Поскважинный потолок — не то же самое, что суммарный по объекту.

    Сто скважин по десять кубов и десять по сто дают один и тот же объект и
    совершенно разные скважины; «ограничение по закачке» может прийти именно в
    этой форме.
    """
    from aios import planner

    steps = _plan(producers=((30.0, 50.0), (10.0, 50.0)),
                  injectors=((200.0, 300.0), (50.0, 300.0)))

    fitted, decisions = planner.enforce(
        steps, _forecast(("P1", "P2")), Norms.from_file(),
        _limits(max_well_liquid_m3_per_day=20.0, max_well_injection_m3_per_day=120.0),
    )

    assert [c.liquid_rate for c in fitted[0].producers] == [20.0, 10.0]
    assert [c.rate for c in fitted[0].injectors] == [120.0, 50.0]
    assert {d.well for d in decisions} == {"P1", "I1"}


def test_an_addressed_cap_overrides_the_cap_by_kind():
    """Адресная карта сильнее потолка по типу: у скважины свой паспорт насоса."""
    from aios import planner

    steps = _plan(producers=((30.0, 50.0), (30.0, 50.0)), injectors=())

    fitted, _decisions = planner.enforce(
        steps, _forecast(("P1", "P2")), Norms.from_file(),
        _limits(max_well_liquid_m3_per_day=25.0, well_caps_m3_per_day={"P2": 5.0}),
    )

    assert [c.liquid_rate for c in fitted[0].producers] == [25.0, 5.0]


def test_forced_shut_wells_are_closed_whatever_the_cycle_proposed():
    """Самая жёсткая форма ограничения по фонду: «эти скважины не трогать»."""
    from aios import planner

    steps = _plan(producers=((30.0, 50.0), (30.0, 50.0)), injectors=((200.0, 300.0),))

    fitted, decisions = planner.enforce(
        steps, _forecast(("P1", "P2")), Norms.from_file(),
        _limits(forced_shut_wells=frozenset({"P2", "I1"})),
    )

    assert [c.status for c in fitted[0].producers] == ["OPEN", "SHUT"]
    assert [c.status for c in fitted[0].injectors] == ["SHUT"]
    assert {d.well for d in decisions} == {"P2", "I1"}
    assert all("обязательн" in d.resource for d in decisions)


def test_the_active_fund_limit_closes_the_worst_by_margin():
    """Предел числа действующих закрывает худших по марже, а не первых попавшихся.

    Правило то же, что при нехватке мощности по жидкости: две разные очереди
    «кого резать первым» разъехались бы молча.
    """
    from aios import planner

    steps = _plan(producers=((30.0, 50.0), (30.0, 50.0)), injectors=())
    # P2 поднимает ту же жидкость почти без нефти — она и есть худшая
    forecast = (_forecast(("P1",), liquid=900.0, oil=450.0)
                + _forecast(("P2",), liquid=900.0, oil=5.0))

    fitted, decisions = planner.enforce(
        steps, forecast, Norms.from_file(), _limits(max_active_wells=1),
    )

    assert [c.status for c in fitted[0].producers] == ["OPEN", "SHUT"]
    assert decisions[0].well == "P2"
    assert "предела 1" in decisions[0].reason


def test_the_active_fund_limit_needs_norms_to_choose():
    """Без нормативов выбирать не из чего — это отказ, а не случайный выбор."""
    from aios import planner

    with pytest.raises(ValueError, match="нужны нормативы"):
        planner.enforce(_plan(), _forecast(), None, _limits(max_active_wells=1))


def test_the_annual_injection_volume_is_not_the_daily_rate():
    """Годовой объём — отдельная форма: суточный предел его не выражает.

    Закачка неравномерна по месяцам, и год может выйти за норму, ни разу не
    превысив суточного потолка. Отраслевая норма нормирует именно объём.
    """
    from aios import planner

    steps = _plan(injectors=((200.0, 300.0),), months=12)
    asked = sum(200.0 * 30 for _ in range(12))

    fitted, decisions = planner.enforce(
        steps, _forecast(), Norms.from_file(),
        _limits(max_field_injection_m3_per_year=asked / 2),
    )

    got = sum(c.rate * 30 for step in fitted for c in step.injectors)
    assert got < asked
    assert decisions and "годовой объём" in decisions[0].resource
    assert "2014" in decisions[0].resource


def test_uncontrollable_wells_keep_the_reference_regime():
    """«Разрешено управлять этими» означает, что остальные остаются как были.

    Форма живёт в цикле, а не в планировщике: чтобы вернуть скважину к
    исходному режиму, нужно опорное расписание, а планировщик получает уже
    изменённые шаги.
    """
    from aios.loop import _restore_uncontrolled

    origin = _plan(producers=((30.0, 50.0), (30.0, 50.0)), injectors=())
    changed = _plan(producers=((45.0, 50.0), (45.0, 50.0)), injectors=())

    restored = _restore_uncontrolled(origin, changed, frozenset({"P1"}))

    assert [c.liquid_rate for c in restored[0].producers] == [45.0, 30.0]
def test_an_exhausted_liquid_budget_closes_injection_too():
    """Год без добычи не должен оставаться годом с закачкой.

    Объём за горизонт раздаётся ранним годам (Р-147), поздним достаётся ноль, и
    всем добывающим ставится SHUT. Закачку при этом не трогали: вода шла в пласт,
    OPEX тратился, выручки не было (ревизия 03.09, 2.1 #6).
    """
    from aios import planner
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep
    from aios.margin import Norms

    steps = [ScheduleStep(
        date(2014 + index // 12, index % 12 + 1, 1),
        (ProducerControl("P1", "OPEN", "LRAT", 100.0, 50.0),),
        (InjectorControl("I1", "WATER", "OPEN", "RATE", 150.0, 300.0),),
    ) for index in range(24)]

    # объёма хватает меньше чем на год: второй год получит ноль
    fitted, decisions = planner.enforce(
        steps, _forecast(months=3), Norms.from_file(),
        _limits(max_field_liquid_m3_total=100.0 * 30 * 6))

    second = [step for step in fitted if step.date.year == 2015]
    assert second, "второй год в расписании есть"
    assert all(c.status == "SHUT" for step in second for c in step.producers)
    assert all(c.status == "SHUT" for step in second for c in step.injectors), (
        "закачка обязана закрыться вместе с отбором"
    )
    assert any("компенсировать нечего" in d.reason for d in decisions)


def test_the_active_wells_limit_closes_producers_not_injectors():
    """Предел числа действующих скважин рубил закачку первой.

    Удельная маржа определена через поднятый объём, а у нагнетательной объёма
    нет — `_margin_per_liquid` возвращает для неё суммарную маржу в млн руб.,
    величину другого масштаба. В общем сорте нагнетательная всегда оказывалась
    худшей, и предел закрывал именно её, хотя срез закачки на Модели Z
    обрушивает нефть вчетверо (ревизия 03.09, 2.1 #5).
    """
    from aios import planner
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep
    from aios.margin import Norms

    steps = [ScheduleStep(
        date(2014, month, 1),
        (ProducerControl("P1", "OPEN", "LRAT", 30.0, 50.0),
         ProducerControl("P2", "OPEN", "LRAT", 30.0, 50.0)),
        (InjectorControl("I1", "WATER", "OPEN", "RATE", 200.0, 300.0),),
    ) for month in (1, 2, 3)]
    forecast = _forecast(wells=("P1", "P2"))

    fitted, decisions = planner.enforce(
        steps, forecast, Norms.from_file(), _limits(max_active_wells=2))

    injecting = [c.well for step in fitted for c in step.injectors if c.status == "OPEN"]
    assert injecting, "закачка не должна закрываться, пока есть добывающие"
    assert any(d.resource == "предел числа действующих скважин" for d in decisions)


def test_the_active_wells_limit_reaches_injectors_when_producers_run_out():
    """Если добывающих не хватило — закрывается самая мелкая закачка."""
    from aios import planner
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep
    from aios.margin import Norms

    steps = [ScheduleStep(
        date(2014, month, 1),
        (ProducerControl("P1", "OPEN", "LRAT", 30.0, 50.0),),
        (InjectorControl("I1", "WATER", "OPEN", "RATE", 50.0, 300.0),
         InjectorControl("I2", "WATER", "OPEN", "RATE", 300.0, 300.0)),
    ) for month in (1, 2, 3)]

    fitted, _ = planner.enforce(
        steps, _forecast(), Norms.from_file(), _limits(max_active_wells=1))

    still = {c.well for step in fitted for c in step.injectors if c.status == "OPEN"}
    assert "I2" not in still or "I1" not in still
    assert "I1" not in still, "мелкая закачка уходит раньше крупной"


def test_uncontrollable_well_returns_to_reference_on_every_step():
    """Опора — расписание-дельты: скважина упомянута раз, действует до конца.

    Дефект ревизии 03.09 (2.1 #1): восстановление искало опорную запись по дате
    в **дельтовом** расписании, и на шагах, где скважину не упоминали, оставляло
    режим планировщика. То есть «управлять нельзя» действовало один месяц из
    трёх, и увидеть это на тесте с полным фондом на каждом шаге нельзя.
    """
    from aios.deck import ProducerControl, ScheduleStep
    from aios.loop import _restore_uncontrolled

    origin = [
        ScheduleStep(date(2014, 1, 1),
                     (ProducerControl("P1", "OPEN", "LRAT", 30.0, 50.0),
                      ProducerControl("P2", "OPEN", "LRAT", 30.0, 50.0)), ()),
        ScheduleStep(date(2014, 2, 1), (), ()),
        ScheduleStep(date(2014, 3, 1), (), ()),
    ]
    changed = _plan(producers=((45.0, 50.0), (45.0, 50.0)), injectors=())

    restored = _restore_uncontrolled(origin, changed, frozenset({"P1"}))

    assert [[c.liquid_rate for c in step.producers] for step in restored] == [
        [45.0, 30.0], [45.0, 30.0], [45.0, 30.0]
    ]


def test_uncontrollable_well_keeps_its_shut_status_from_the_reference():
    """Возврат к опоре возвращает и статус, а не только уставку."""
    from aios.deck import ProducerControl, ScheduleStep
    from aios.loop import _restore_uncontrolled

    origin = [
        ScheduleStep(date(2014, 1, 1),
                     (ProducerControl("P1", "OPEN", "LRAT", 30.0, 50.0),
                      ProducerControl("P2", "SHUT", "LRAT", None, 50.0)), ()),
        ScheduleStep(date(2014, 2, 1), (), ()),
        ScheduleStep(date(2014, 3, 1), (), ()),
    ]
    changed = _plan(producers=((45.0, 50.0), (45.0, 50.0)), injectors=())

    restored = _restore_uncontrolled(origin, changed, frozenset({"P1"}))

    assert [step.producers[1].status for step in restored] == ["SHUT", "SHUT", "SHUT"]


def test_restoring_a_well_that_breaks_an_object_limit_is_reported():
    """Возврат к опоре идёт после планировщика и может сломать его срезку.

    Это не лечится порядком: «управлять нельзя» и «объектный предел» могут
    противоречить друг другу, и тогда кейс противоречив. Молчать нельзя —
    условие обязано быть названо нарушенным (`case_problems`, А2 ревизии).
    """
    from aios import casecheck, loop

    limits = _limits(max_field_liquid_m3_per_day=60.0,
                     controllable_wells=frozenset({"P1"}))
    steps = _plan(producers=((30.0, 50.0), (30.0, 50.0)), injectors=())
    forecast = _forecast(wells=("P1", "P2"))

    left = casecheck.violations(casecheck.check(steps, forecast, limits))
    assert not left, "опора сама предел не нарушает"

    broken = _plan(producers=((45.0, 50.0), (45.0, 50.0)), injectors=())
    reported = loop._case_problems(broken, forecast, limits, None, [])

    assert any("max_field_liquid_m3_per_day" in line for line in reported)


def _switching(states, injecting=()):
    """Расписание одной скважины, где статус меняется от шага к шагу.

    `states` — статус P1 по месяцам; `injecting` — месяцы, в которых у той же
    скважины появляется запись закачки (так выглядит перевод под закачку:
    добыча закрывается, `WCONINJE` дописывается).
    """
    from aios.deck import InjectorControl, ProducerControl, ScheduleStep

    steps = []
    for index, status in enumerate(states, start=1):
        rate = 30.0 if status == "OPEN" else None
        producers = (ProducerControl("P1", status, "LRAT", rate, 50.0),)
        injectors = ((InjectorControl("P1", "WATER", "OPEN", "RATE", 30.0, 300.0),)
                     if index in injecting else ())
        steps.append(ScheduleStep(date(2014, index, 1), producers, injectors))
    return steps


def test_the_event_quota_is_enforced_and_not_only_parsed():
    """`max_events_per_year` до 30.08 разбирался конфигурацией и ничего не делал.

    Методика оценивает мероприятия поштучно (1,0 млн руб. за остановку и за
    пуск), поэтому «не более N мероприятий в год» — естественная форма
    ограничения по фонду. Квота расходуется по календарю: первые N мероприятий
    года проходят, остальные отменяются, и скважина остаётся в том состоянии,
    в котором была.
    """
    from aios import planner

    # стоп, пуск, стоп — три мероприятия за 2014 год
    steps = _switching(["OPEN", "SHUT", "OPEN", "SHUT"])

    fitted, decisions = planner.enforce(
        steps, _forecast(), Norms.from_file(), _limits(max_events_per_year=1),
    )

    # Состояние фонда, а не состав шага: с 02.09 отмена мероприятия снимает
    # запись с шага, а не подставляет запись предыдущего. В деке отсутствие
    # записи означает «режим не менялся», и прежняя редакция на настоящем
    # расписании (шаг несёт только изменившиеся скважины) не срабатывала вовсе.
    assert [state["P1"] for state in planner.fund_states(fitted)] == [
        (True, False), (False, False), (False, False), (False, False)]
    assert len(decisions) == 1
    assert decisions[0].well == "P1"
    assert "2014" in decisions[0].resource
    assert "пуск" in decisions[0].reason and "предел 1" in decisions[0].reason


def test_the_event_quota_counts_a_conversion_once():
    """Перевод под закачку — одно мероприятие, а не остановка плюс пуск.

    В расписании он выглядит двумя записями: добыча закрывается, `WCONINJE`
    дописывается. Считать это за два мероприятия значило бы вдвое сузить квоту.
    """
    from aios import planner

    steps = _switching(["OPEN", "SHUT", "SHUT"], injecting=(2, 3))

    passes, none = planner.enforce(
        steps, _forecast(), Norms.from_file(), _limits(max_events_per_year=1),
    )
    assert none == []
    assert [len(step.injectors) for step in passes] == [0, 1, 1]

    # квота уже израсходована пуском — перевод отменён целиком: и добыча
    # возвращена, и запись закачки, которой до перевода не было, снята
    spent = _switching(["SHUT", "OPEN", "SHUT"], injecting=(3,))
    reverted, decisions = planner.enforce(
        spent, _forecast(), Norms.from_file(), _limits(max_events_per_year=1),
    )
    # та же смена системы отсчёта: смотрим состояние фонда, а не записи
    assert [state["P1"] for state in planner.fund_states(reverted)] == [
        (False, False), (True, False), (True, False)]
    assert [len(step.injectors) for step in reverted] == [0, 0, 0]
    assert len(decisions) == 1
    assert "перевод под закачку" in decisions[0].reason


def test_without_the_event_key_nothing_changes():
    """Умолчание ручки — «ограничения нет»: ни одного решения и ни одной правки."""
    from aios import planner

    steps = _switching(["OPEN", "SHUT", "OPEN"])

    fitted, decisions = planner.enforce(
        steps, _forecast(), Norms.from_file(), _limits(),
    )

    assert decisions == []
    assert [step.producers[0].status for step in fitted] == ["OPEN", "SHUT", "OPEN"]
